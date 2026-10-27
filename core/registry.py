# -*- coding: utf-8 -*-
"""
Módulo de Registro e Dicionário de Tags XML (core.registry)
Fonte Única da Verdade (Single Source of Truth) baseada em dic-tags-api.xlsx.
Fornece funções de carga, busca, auditoria e validação cruzada de schemas.
"""

import os
import re
import openpyxl
from core.config import EXCEL_DICT_PATH


class TagRegistry:
    """Gerenciador do Dicionário de Tags XML."""

    def __init__(self, excel_path: str = EXCEL_DICT_PATH):
        self.excel_path = excel_path
        self._tags = []
        self._tags_by_name = {}
        self._loaded = False

    def load(self, force_reload: bool = False) -> list:
        """Carrega e valida as tags da planilha oficial Excel."""
        if self._loaded and not force_reload:
            return self._tags

        if not os.path.exists(self.excel_path):
            raise FileNotFoundError(f"Planilha de tags não encontrada em: {self.excel_path}")

        wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        ws = wb.active

        tags = []
        tags_by_name = {}

        for r in range(2, ws.max_row + 1):
            tag_cell = ws.cell(row=r, column=1).value
            conc_cell = ws.cell(row=r, column=2).value
            regr_cell = ws.cell(row=r, column=3).value

            if tag_cell is None:
                continue

            raw_tag = str(tag_cell).strip()
            # Garante formato <nomeDaTag>
            if not raw_tag.startswith("<"):
                raw_tag = f"<{raw_tag}"
            if not raw_tag.endswith(">"):
                raw_tag = f"{raw_tag}>"

            conceito = str(conc_cell).strip() if conc_cell is not None else ""
            regra = str(regr_cell).strip() if regr_cell is not None else ""

            tag_item = {
                "name": raw_tag,
                "clean_name": raw_tag.strip("<>").strip(),
                "conceito": conceito,
                "regra": regra
            }

            tags.append(tag_item)
            tags_by_name[tag_item["name"]] = tag_item
            tags_by_name[tag_item["clean_name"]] = tag_item

        self._tags = tags
        self._tags_by_name = tags_by_name
        self._loaded = True
        return self._tags

    def get_all_tags(self) -> list:
        """Retorna a lista completa de tags formatada para o frontend."""
        if not self._loaded:
            self.load()
        return [{"name": t["name"], "conceito": t["conceito"], "regra": t["regra"]} for t in self._tags]

    def get_tag_count(self) -> int:
        """Retorna o total de tags registradas."""
        if not self._loaded:
            self.load()
        return len(self._tags)

    def get_tag(self, tag_name: str) -> dict:
        """Busca uma tag por seu nome com ou sem delimitadores <>."""
        if not self._loaded:
            self.load()
        clean = tag_name.strip("<>").strip()
        return self._tags_by_name.get(clean) or self._tags_by_name.get(f"<{clean}>")

    def has_tag(self, tag_name: str) -> bool:
        """Verifica se a tag existe no dicionário."""
        return self.get_tag(tag_name) is not None

    def extract_tags_from_xml(self, xml_text: str) -> set:
        """Extrai todos os nomes de tags presentes em uma string XML."""
        if not xml_text:
            return set()
        
        # Expressão para capturar nomes de tags XML (ignorando prefixos de namespace e tags de fechamento)
        matches = re.findall(r'<(?:\/?(?:[a-zA-Z0-9_\-]+:)?)([a-zA-Z0-9_]+)(?:[^>]*)>', xml_text)
        
        # Filtros de tags estruturais padrão SOAP / XML e wrappers de subtipos
        ignorar = {
            "xml", "soap", "soapenv", "soap-env", "Envelope", "Header", "Body", "Fault",
            "return", "resposta", "int", "grp", "ns2", "pag",
            "itemDePagamento"
        }
        
        return {m for m in matches if m not in ignorar}

    def audit_xml(self, xml_text: str) -> dict:
        """
        Audita um XML contra o dicionário de tags.
        Retorna tags reconhecidas e tags que não constam no dicionário.
        """
        xml_tags = self.extract_tags_from_xml(xml_text)
        reconhecidas = set()
        nao_registradas = set()

        for tag in xml_tags:
            if self.has_tag(tag):
                reconhecidas.add(tag)
            else:
                nao_registradas.add(tag)

        return {
            "total_tags_xml": len(xml_tags),
            "reconhecidas": sorted(list(reconhecidas)),
            "nao_registradas": sorted(list(nao_registradas))
        }


# Instância singleton padrão
default_registry = TagRegistry()
